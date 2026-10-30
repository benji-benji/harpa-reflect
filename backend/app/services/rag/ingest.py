import re
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

import fitz  # PyMuPDF
import tiktoken
from openai import OpenAI

from app.core.config import settings
from app.core.database import get_supabase


ENCODING = tiktoken.get_encoding("cl100k_base")
MAX_TOKENS = 800
EMBED_MODEL = "text-embedding-3-small"
EMBED_BATCH = 128


@dataclass
class Chunk:
    page_number: int
    chunk_index: int
    content: str
    token_count: int


def _count_tokens(text: str) -> int:
    return len(ENCODING.encode(text))


def _split_sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]


def _sanitize(text: str) -> str:
    # Postgres text columns reject the U+0000 null byte
    return text.replace("\x00", "")


def _extract_pages(pdf_path: Path) -> list[tuple[int, str]]:
    pages = []
    with fitz.open(pdf_path) as doc:
        for page in doc:
            text = _sanitize(page.get_text())
            if text.strip():
                pages.append((page.number + 1, text))
    return pages


def _extract_docx(docx_path: Path) -> list[tuple[int, str]]:
    from docx import Document
    doc = Document(str(docx_path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = (cell.text or "").strip()
                if t:
                    parts.append(t)
    text = _sanitize("\n".join(parts)).strip()
    return [(1, text)] if text else []


def _extract(path: Path) -> list[tuple[int, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pages(path)
    if suffix == ".docx":
        return _extract_docx(path)
    return []


def _chunk_pages(pages: list[tuple[int, str]]) -> list[Chunk]:
    chunks: list[Chunk] = []
    for page_number, raw in pages:
        text = re.sub(r"\n{3,}", "\n\n", raw).strip()
        if not text:
            continue
        tokens = _count_tokens(text)
        if tokens <= MAX_TOKENS:
            chunks.append(Chunk(page_number, 0, text, tokens))
            continue
        sentences = _split_sentences(text)
        current: list[str] = []
        current_tokens = 0
        idx = 1
        for sentence in sentences:
            st = _count_tokens(sentence)
            if current_tokens + st > MAX_TOKENS and current:
                content = " ".join(current)
                chunks.append(Chunk(page_number, idx, content, _count_tokens(content)))
                idx += 1
                current = current[-1:]
                current_tokens = _count_tokens(current[0]) if current else 0
            current.append(sentence)
            current_tokens += st
        if current:
            content = " ".join(current)
            chunks.append(Chunk(page_number, idx, content, _count_tokens(content)))
    return chunks


def _embed(texts: list[str]) -> list[list[float]]:
    client = OpenAI(api_key=settings.openai_api_key)
    out: list[list[float]] = []
    for i in range(0, len(texts), EMBED_BATCH):
        batch = texts[i : i + EMBED_BATCH]
        response = client.embeddings.create(input=batch, model=EMBED_MODEL)
        out.extend(item.embedding for item in response.data)
    return out


def _purge_empty_documents() -> int:
    sb = get_supabase()
    docs = sb.table("theory_documents").select("id,filename").execute()
    removed = 0
    for d in docs.data:
        chunks = sb.table("theory_chunks").select("id", count="exact").eq("document_id", d["id"]).limit(1).execute()
        if (chunks.count or 0) == 0:
            sb.table("theory_documents").delete().eq("id", d["id"]).execute()
            print(f"purge {d['filename']} (0 chunks)")
            removed += 1
    return removed


SUPPORTED_SUFFIXES = {".pdf", ".docx"}


def ingest_directory(pdf_dir: Path, recursive: bool = False) -> dict:
    sb = get_supabase()
    _purge_empty_documents()
    existing = sb.table("theory_documents").select("filename").execute()
    already = {row["filename"] for row in existing.data}

    iterator = pdf_dir.rglob("*") if recursive else pdf_dir.iterdir()
    pdfs = sorted(p for p in iterator if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES)
    stats = {"total": len(pdfs), "ingested": 0, "skipped": 0, "chunks": 0}

    for pdf_path in pdfs:
        if pdf_path.name in already:
            stats["skipped"] += 1
            print(f"skip  {pdf_path.name}")
            continue

        pages = _extract(pdf_path)
        if not pages:
            print(f"empty {pdf_path.name}")
            continue

        chunks = _chunk_pages(pages)
        if not chunks:
            print(f"nochk {pdf_path.name}")
            continue

        print(f"embed {pdf_path.name}  pages={len(pages)}  chunks={len(chunks)}")
        embeddings = _embed([c.content for c in chunks])

        doc_id = str(uuid4())
        sb.table("theory_documents").insert({
            "id": doc_id,
            "filename": pdf_path.name,
            "title": pdf_path.stem,
            "total_pages": len(pages),
        }).execute()

        rows = [
            {
                "id": str(uuid4()),
                "document_id": doc_id,
                "page_number": c.page_number,
                "chunk_index": c.chunk_index,
                "content": c.content,
                "embedding": "[" + ",".join(f"{x:.7f}" for x in emb) + "]",
                "token_count": c.token_count,
            }
            for c, emb in zip(chunks, embeddings)
        ]
        for i in range(0, len(rows), 100):
            sb.table("theory_chunks").insert(rows[i : i + 100]).execute()

        stats["ingested"] += 1
        stats["chunks"] += len(chunks)

    return stats
